"""论文预览与导出版式回归测试。"""

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.utils.paper_formatting import (
    create_reference_docx,
    format_exported_docx,
    normalize_paper_numbering,
)


class TestPaperFormatting(unittest.TestCase):
    """验证预览、DOCX 与 PDF 共用的编号和页面参数。"""

    def test_numbering_is_continuous_and_skips_code_fences(self):
        """公式、图、表连续编号，代码示例不参与编号。"""
        markdown = r"""
$$
x = 1 \qquad (8)
$$

```markdown
$$
example = 1
$$
![示例](ignored.png)
```

$$
y = 2
$$

![旧图题](first.png)
![图 9  第二张图](second.png)

**表9 指标对比**
"""

        normalized = normalize_paper_numbering(markdown)

        self.assertIn(r"x = 1", normalized)
        self.assertNotIn(r"\qquad (8)", normalized)
        self.assertIn(r"\tag{1}", normalized)
        self.assertIn(r"\tag{2}", normalized)
        self.assertIn("![图 1  旧图题](first.png)", normalized)
        self.assertIn("![图 2  第二张图](second.png)", normalized)
        self.assertIn("**表 1  指标对比**", normalized)
        self.assertIn("![示例](ignored.png)", normalized)

        export_source = normalize_paper_numbering(
            markdown,
            include_equation_tags=False,
            explicit_caption_numbers=False,
        )
        self.assertNotIn(r"\tag{", export_source)
        self.assertIn("Table: 指标对比", export_source)
        self.assertIn("![旧图题](first.png)", export_source)

    def test_reference_docx_matches_official_page_and_body_style(self):
        """参考 DOCX 必须使用 A4、指定页边距和固定 24 磅正文。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            path = create_reference_docx(Path(temp_dir) / "reference.docx")
            document = Document(path)
            section = document.sections[0]
            normal = document.styles["Normal"]

            self.assertAlmostEqual(section.page_width.cm, 21.0, places=1)
            self.assertAlmostEqual(section.page_height.cm, 29.7, places=1)
            self.assertAlmostEqual(section.top_margin.cm, 2.54, places=1)
            self.assertAlmostEqual(section.left_margin.cm, 3.17, places=1)
            self.assertEqual(
                normal.paragraph_format.line_spacing_rule,
                WD_LINE_SPACING.EXACTLY,
            )
            self.assertAlmostEqual(
                normal.paragraph_format.line_spacing.pt,
                24,
                places=1,
            )
            self.assertAlmostEqual(
                normal.paragraph_format.first_line_indent.pt,
                24,
                places=1,
            )

    def test_docx_postprocessing_formats_body_equation_and_three_line_table(self):
        """二次排版必须覆盖正文、公式编号和三线表。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "paper.docx"
            document = Document()
            document.add_heading("统计建模论文标题", level=1)
            document.add_paragraph("这是正文段落。")
            equation = document.add_paragraph()
            office_math = OxmlElement("m:oMath")
            math_run = OxmlElement("m:r")
            math_text = OxmlElement("m:t")
            math_text.text = "x=1"
            math_run.append(math_text)
            office_math.append(math_run)
            equation._p.append(office_math)
            document.styles.add_style(
                "Table Caption",
                WD_STYLE_TYPE.PARAGRAPH,
            )
            document.add_paragraph(
                "指标对比",
                style="Table Caption",
            )
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "指标"
            table.cell(0, 1).text = "数值"
            table.cell(1, 0).text = "RMSE"
            table.cell(1, 1).text = "1.2"
            document.save(path)

            format_exported_docx(path)
            formatted = Document(path)
            formatted_body = formatted.paragraphs[1]
            formatted_equation = formatted.paragraphs[2]
            formatted_caption = formatted.paragraphs[3]
            borders = formatted.tables[0]._tbl.tblPr.find(
                qn("w:tblBorders")
            )

            self.assertEqual(
                formatted_body.alignment,
                WD_ALIGN_PARAGRAPH.JUSTIFY,
            )
            self.assertEqual(
                formatted_body.paragraph_format.line_spacing_rule,
                WD_LINE_SPACING.EXACTLY,
            )
            self.assertEqual(
                formatted_body.paragraph_format.first_line_indent,
                Pt(24),
            )
            self.assertIn("（1）", formatted_equation.text)
            self.assertEqual(formatted_caption.text, "表 1  指标对比")
            self.assertIsNotNone(borders)
            self.assertEqual(
                borders.find(qn("w:top")).get(qn("w:val")),
                "single",
            )
            self.assertEqual(
                borders.find(qn("w:left")).get(qn("w:val")),
                "nil",
            )
            self.assertAlmostEqual(
                formatted.sections[0].left_margin,
                Cm(3.17),
                delta=200,
            )

    def test_docx_toc_is_dynamic_and_uses_independent_pages(self):
        """DOCX 的摘要、目录与正文必须各自起页，目录使用可更新域。"""
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "paper-with-toc.docx"
            document = Document()
            document.add_paragraph("示例论文", style="Title")
            document.add_heading("摘 要", level=1)
            document.add_paragraph("摘要正文。")
            document.add_heading("目 录", level=1)
            document.add_paragraph("一、问题重述 1")
            document.add_paragraph("1.1 问题背景 1")
            document.add_heading("一、问题重述", level=1)
            document.add_heading("1.1 问题背景", level=2)
            document.add_paragraph("正文。")
            document.save(path)

            format_exported_docx(path)
            formatted = Document(path)
            paragraphs = formatted.paragraphs
            texts = [paragraph.text for paragraph in paragraphs]
            toc_heading = next(
                paragraph
                for paragraph in paragraphs
                if paragraph.text.replace(" ", "") == "目录"
            )
            body_heading = next(
                paragraph
                for paragraph in paragraphs
                if paragraph.text == "一、问题重述"
            )

            self.assertNotIn("一、问题重述 1", texts)
            self.assertIn('TOC \\o "1-3"', formatted._element.xml)
            self.assertTrue(toc_heading.paragraph_format.page_break_before)
            self.assertTrue(body_heading.paragraph_format.page_break_before)
            self.assertEqual(
                toc_heading.alignment,
                WD_ALIGN_PARAGRAPH.CENTER,
            )
            self.assertIn("updateFields", formatted.settings._element.xml)


if __name__ == "__main__":
    unittest.main()
