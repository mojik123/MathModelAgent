"""统计建模论文的统一排版与编号规则。

该模块把网页预览、DOCX 与 PDF 共用的版式参数集中为一份权威配置，
避免不同导出链路分别硬编码后产生字号、行距和编号漂移。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt
from docx.text.paragraph import Paragraph

A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7
MARGIN_VERTICAL_CM = 2.54
MARGIN_HORIZONTAL_CM = 3.17
BODY_FONT_SIZE_PT = 12
BODY_LINE_SPACING_PT = 24
TABLE_FONT_SIZE_PT = 10.5
TITLE_FONT_SIZE_PT = 16
HEADING_1_FONT_SIZE_PT = 15
HEADING_2_FONT_SIZE_PT = 14
HEADING_3_FONT_SIZE_PT = 12
FRONT_HEADING_FONT_SIZE_PT = 14

WESTERN_FONT = "Times New Roman"
BODY_CJK_FONT = "SimSun"
TITLE_CJK_FONT = "FZXiaoBiaoSong-B05S"
HEADING_1_CJK_FONT = "SimHei"
HEADING_2_CJK_FONT = "KaiTi"

_FENCED_CODE_RE = re.compile(r"(```[\s\S]*?```|~~~[\s\S]*?~~~)")
_DISPLAY_MATH_RE = re.compile(
    r"(?P<dollar>\$\$(?P<dollar_body>[\s\S]*?)\$\$)"
    r"|(?P<bracket>\\\[(?P<bracket_body>[\s\S]*?)\\\])"
)
_IMAGE_RE = re.compile(r"!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)")
_TABLE_CAPTION_RE = re.compile(
    r"^(?P<prefix>\s*(?:\*\*)?)表\s*(?:\d+(?:\.\d+)?)?"
    r"[\s　:：、.-]*(?P<caption>.*?)(?P<suffix>(?:\*\*)?\s*)$"
)
_FIGURE_PREFIX_RE = re.compile(
    r"^\s*图\s*(?:\d+(?:\.\d+)?)?[\s　:：、.-]*(?P<caption>.*)$"
)
_NUMBERED_EQUATION_RE = re.compile(r"[（(]\s*\d+\s*[）)]")


def normalize_paper_numbering(
    markdown: str,
    *,
    include_equation_tags: bool = True,
    explicit_caption_numbers: bool = True,
) -> str:
    """为公式、插图和显式表题生成全文连续编号。

    代码围栏中的 Markdown 示例不会参与编号，避免修改附录代码。

    Args:
        markdown: 原始论文 Markdown。

    Returns:
        编号已规范化的 Markdown。
    """
    equation_index = 0
    figure_index = 0
    table_index = 0

    def normalize_text(segment: str) -> str:
        nonlocal equation_index, figure_index, table_index

        def number_equation(match: re.Match[str]) -> str:
            nonlocal equation_index
            body = match.group("dollar_body")
            delimiter = "dollar"
            if body is None:
                body = match.group("bracket_body") or ""
                delimiter = "bracket"

            stripped = body.strip()
            if not stripped or re.search(r"\\begin\{align\*?\}", stripped):
                return match.group(0)

            equation_index += 1
            stripped = re.sub(
                r"\\tag\{[^{}]*\}\s*$",
                "",
                stripped,
            ).rstrip()
            stripped = re.sub(
                r"(?:\\q?quad\s*)?[（(]\s*\d+\s*[）)]\s*$",
                "",
                stripped,
            ).rstrip()
            numbered = (
                f"{stripped}\n\\tag{{{equation_index}}}"
                if include_equation_tags
                else stripped
            )
            if delimiter == "dollar":
                return f"$$\n{numbered}\n$$"
            return f"\\[\n{numbered}\n\\]"

        def number_figure(match: re.Match[str]) -> str:
            nonlocal figure_index
            alt = match.group("alt").strip()
            src = match.group("src")
            caption_match = _FIGURE_PREFIX_RE.match(alt)
            if caption_match:
                caption = caption_match.group("caption").strip()
            else:
                caption = alt
            if not caption or _looks_like_filename(caption):
                caption = "模型结果"
            figure_index += 1
            if explicit_caption_numbers:
                return f"![图 {figure_index}  {caption}]({src})"
            return f"![{caption}]({src})"

        output = _DISPLAY_MATH_RE.sub(number_equation, segment)
        output = _IMAGE_RE.sub(number_figure, output)

        normalized_lines: list[str] = []
        for line in output.splitlines(keepends=True):
            line_ending = "\n" if line.endswith("\n") else ""
            raw_line = line[:-1] if line_ending else line
            caption_match = _TABLE_CAPTION_RE.match(raw_line)
            if not caption_match:
                normalized_lines.append(line)
                continue
            caption = caption_match.group("caption").strip()
            if not caption:
                normalized_lines.append(line)
                continue
            table_index += 1
            if not explicit_caption_numbers:
                normalized_lines.append(f"Table: {caption}{line_ending}")
                continue
            prefix = caption_match.group("prefix")
            suffix = caption_match.group("suffix")
            normalized_lines.append(
                f"{prefix}表 {table_index}  {caption}{suffix}{line_ending}"
            )
        return "".join(normalized_lines)

    parts = _FENCED_CODE_RE.split(markdown)
    return "".join(
        part if _FENCED_CODE_RE.fullmatch(part) else normalize_text(part)
        for part in parts
    )


def create_reference_docx(output_path: str | Path) -> Path:
    """创建供 Pandoc 使用的统计建模论文参考 DOCX。

    Args:
        output_path: 参考文档保存路径。

    Returns:
        保存后的参考文档路径。
    """
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    _configure_document(document)
    document.save(path)
    return path


def format_exported_docx(docx_path: str | Path) -> Path:
    """对 Pandoc 生成的 DOCX 执行严格的二次排版。

    Args:
        docx_path: 待格式化的 DOCX 路径。

    Returns:
        已原位更新的 DOCX 路径。

    Raises:
        FileNotFoundError: DOCX 不存在。
    """
    path = Path(docx_path)
    if not path.is_file():
        raise FileNotFoundError(f"DOCX 不存在: {path}")

    document = Document(path)
    _configure_document(document)
    _normalize_docx_caption_numbers(document)
    _prepare_docx_front_matter(document)

    first_heading_seen = False
    equation_index = 0
    for paragraph in document.paragraphs:
        style_name = (paragraph.style.name or "").lower()
        text = paragraph.text.strip()
        is_equation = _contains_equation(paragraph)
        is_drawing = _contains_drawing(paragraph)

        if ' TOC \\o "1-3"' in paragraph._p.xml:
            _format_toc_field_paragraph(paragraph)
            continue
        if is_equation:
            equation_index += 1
            _format_equation_paragraph(document, paragraph, equation_index)
            continue
        if is_drawing:
            _format_figure_paragraph(paragraph)
            continue
        if _is_caption(text, style_name):
            _format_caption_paragraph(paragraph)
            continue
        if style_name in {"title", "subtitle"}:
            _format_title_paragraph(paragraph)
            first_heading_seen = True
            continue
        if _compact_heading_text(text) in {"摘要", "目录"}:
            _format_front_heading_paragraph(paragraph)
            first_heading_seen = True
            continue
        if style_name.startswith("heading 1"):
            if not first_heading_seen and _looks_like_document_title(text):
                _format_title_paragraph(paragraph)
            else:
                _format_heading_paragraph(
                    paragraph,
                    HEADING_1_CJK_FONT,
                    HEADING_1_FONT_SIZE_PT,
                    bold=True,
                )
            first_heading_seen = True
            continue
        if style_name.startswith("heading 2"):
            _format_heading_paragraph(
                paragraph,
                HEADING_2_CJK_FONT,
                HEADING_2_FONT_SIZE_PT,
            )
            continue
        if style_name.startswith("heading 3"):
            _format_heading_paragraph(
                paragraph,
                BODY_CJK_FONT,
                HEADING_3_FONT_SIZE_PT,
                bold=True,
            )
            continue
        if style_name.startswith("heading"):
            _format_heading_paragraph(
                paragraph,
                BODY_CJK_FONT,
                BODY_FONT_SIZE_PT,
            )
            continue
        _format_body_paragraph(paragraph, is_list="list" in style_name)

    for table in document.tables:
        _format_table(table)

    document.save(path)
    return path


def _configure_document(document: DocumentObject) -> None:
    """设置页面、样式和页脚。"""
    for section in document.sections:
        section.start_type = WD_SECTION.NEW_PAGE
        section.page_width = Cm(A4_WIDTH_CM)
        section.page_height = Cm(A4_HEIGHT_CM)
        section.top_margin = Cm(MARGIN_VERTICAL_CM)
        section.bottom_margin = Cm(MARGIN_VERTICAL_CM)
        section.left_margin = Cm(MARGIN_HORIZONTAL_CM)
        section.right_margin = Cm(MARGIN_HORIZONTAL_CM)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)
        _add_page_number(section.footer.paragraphs[0])

    _configure_style(
        document,
        "Normal",
        BODY_CJK_FONT,
        BODY_FONT_SIZE_PT,
        first_line_indent=True,
    )
    _configure_style(
        document,
        "Body Text",
        BODY_CJK_FONT,
        BODY_FONT_SIZE_PT,
        first_line_indent=True,
    )
    _configure_style(
        document,
        "Title",
        TITLE_CJK_FONT,
        TITLE_FONT_SIZE_PT,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _configure_style(
        document,
        "Subtitle",
        BODY_CJK_FONT,
        BODY_FONT_SIZE_PT,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _configure_style(
        document,
        "Heading 1",
        HEADING_1_CJK_FONT,
        HEADING_1_FONT_SIZE_PT,
        bold=True,
    )
    _configure_style(
        document,
        "Heading 2",
        HEADING_2_CJK_FONT,
        HEADING_2_FONT_SIZE_PT,
    )
    _configure_style(
        document,
        "Heading 3",
        BODY_CJK_FONT,
        HEADING_3_FONT_SIZE_PT,
        bold=True,
    )
    _configure_style(
        document,
        "Heading 4",
        BODY_CJK_FONT,
        BODY_FONT_SIZE_PT,
    )
    _configure_style(
        document,
        "Caption",
        BODY_CJK_FONT,
        BODY_FONT_SIZE_PT,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
    )
    for caption_style_name in ("Table Caption", "Image Caption"):
        if caption_style_name not in document.styles:
            document.styles.add_style(
                caption_style_name,
                WD_STYLE_TYPE.PARAGRAPH,
            )
        _configure_style(
            document,
            caption_style_name,
            BODY_CJK_FONT,
            BODY_FONT_SIZE_PT,
            alignment=WD_ALIGN_PARAGRAPH.CENTER,
        )
    if "Table Text" not in document.styles:
        document.styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    _configure_style(
        document,
        "Table Text",
        BODY_CJK_FONT,
        TABLE_FONT_SIZE_PT,
        line_spacing_pt=None,
    )


def _configure_style(
    document: DocumentObject,
    name: str,
    cjk_font: str,
    size_pt: float,
    *,
    bold: bool = False,
    alignment: WD_ALIGN_PARAGRAPH | None = None,
    first_line_indent: bool = False,
    line_spacing_pt: float | None = BODY_LINE_SPACING_PT,
) -> None:
    """配置一个 Word 样式。"""
    style = document.styles[name]
    style.font.name = WESTERN_FONT
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    _set_style_east_asia_font(style, cjk_font)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    if line_spacing_pt is None:
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    else:
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        style.paragraph_format.line_spacing = Pt(line_spacing_pt)
    style.paragraph_format.first_line_indent = (
        Pt(BODY_FONT_SIZE_PT * 2) if first_line_indent else Pt(0)
    )
    if alignment is not None:
        style.paragraph_format.alignment = alignment


def _set_style_east_asia_font(style, font_name: str) -> None:
    """为样式设置中文字体。"""
    style.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"),
        font_name,
    )


def _set_run_font(
    run,
    cjk_font: str,
    size_pt: float,
    *,
    bold: bool | None = None,
) -> None:
    """设置文本 run 的中西文字体。"""
    run.font.name = WESTERN_FONT
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn("w:eastAsia"),
        cjk_font,
    )


def _base_paragraph_format(paragraph: Paragraph) -> None:
    """应用固定 24 磅、段前后 0 的基础段落格式。"""
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(BODY_LINE_SPACING_PT)


def _format_body_paragraph(paragraph: Paragraph, *, is_list: bool) -> None:
    """格式化正文段落。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.first_line_indent = (
        Pt(0) if is_list else Pt(BODY_FONT_SIZE_PT * 2)
    )
    for run in paragraph.runs:
        _set_run_font(run, BODY_CJK_FONT, BODY_FONT_SIZE_PT)


def _format_title_paragraph(paragraph: Paragraph) -> None:
    """格式化论文总标题。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    for run in paragraph.runs:
        _set_run_font(
            run,
            TITLE_CJK_FONT,
            TITLE_FONT_SIZE_PT,
            bold=False,
        )


def _format_front_heading_paragraph(paragraph: Paragraph) -> None:
    """格式化摘要、目录等需要居中显示的前置标题。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.page_break_before = True
    for run in paragraph.runs:
        _set_run_font(
            run,
            HEADING_1_CJK_FONT,
            FRONT_HEADING_FONT_SIZE_PT,
            bold=True,
        )


def _format_toc_field_paragraph(paragraph: Paragraph) -> None:
    """格式化 Word 目录域的占位段落，避免继承正文首行缩进。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    for run in paragraph.runs:
        _set_run_font(run, BODY_CJK_FONT, BODY_FONT_SIZE_PT)


def _format_heading_paragraph(
    paragraph: Paragraph,
    cjk_font: str,
    size_pt: float,
    *,
    bold: bool = False,
) -> None:
    """格式化章节标题。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    for run in paragraph.runs:
        _set_run_font(run, cjk_font, size_pt, bold=bold)


def _format_caption_paragraph(paragraph: Paragraph) -> None:
    """格式化图题与表题。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    for run in paragraph.runs:
        _set_run_font(run, BODY_CJK_FONT, BODY_FONT_SIZE_PT, bold=False)


def _normalize_docx_caption_numbers(document: DocumentObject) -> None:
    """为 Pandoc 的图题和表题补充 Word 中可见的连续编号。"""
    figure_index = 0
    table_index = 0
    for paragraph in document.paragraphs:
        style_id = _paragraph_style_id(paragraph)
        text = paragraph.text.strip()
        if style_id == "TableCaption" or re.match(r"^\s*表\s*\d+", text):
            table_index += 1
            caption = re.sub(
                r"^\s*表\s*\d+(?:\.\d+)?[\s　:：、.-]*",
                "",
                text,
            ).strip()
            paragraph.text = f"表 {table_index}  {caption or '数据结果'}"
        elif style_id == "ImageCaption" or re.match(r"^\s*图\s*\d+", text):
            figure_index += 1
            caption = re.sub(
                r"^\s*图\s*\d+(?:\.\d+)?[\s　:：、.-]*",
                "",
                text,
            ).strip()
            paragraph.text = f"图 {figure_index}  {caption or '模型结果'}"


def _paragraph_style_id(paragraph: Paragraph) -> str:
    """读取段落底层样式 ID，兼容参考 DOCX 中暂未注册的样式。"""
    properties = paragraph._p.pPr
    if properties is None or properties.pStyle is None:
        return ""
    return properties.pStyle.val or ""


def _format_figure_paragraph(paragraph: Paragraph) -> None:
    """让图片所在段落居中且不缩进。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)


def _format_equation_paragraph(
    document: DocumentObject,
    paragraph: Paragraph,
    equation_index: int,
) -> None:
    """将公式居中并把编号放到正文版心右侧。"""
    _base_paragraph_format(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    if _NUMBERED_EQUATION_RE.search(paragraph.text):
        return

    section = document.sections[0]
    content_width = section.page_width - section.left_margin - section.right_margin
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(
        Emu(content_width // 2),
        WD_TAB_ALIGNMENT.CENTER,
        WD_TAB_LEADER.SPACES,
    )
    tab_stops.add_tab_stop(
        Emu(content_width),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.SPACES,
    )

    paragraph_element = paragraph._p
    first_content_index = 1 if paragraph_element.pPr is not None else 0
    tab_run = OxmlElement("w:r")
    tab_run.append(OxmlElement("w:tab"))
    paragraph_element.insert(first_content_index, tab_run)
    number_run = paragraph.add_run(f"\t（{equation_index}）")
    _set_run_font(number_run, BODY_CJK_FONT, BODY_FONT_SIZE_PT)


def _format_table(table) -> None:
    """应用三线表、单倍行距与表内字号。"""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    _set_table_borders(table)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell, top=60, start=80, bottom=60, end=80)
            if row_index == 0:
                _set_cell_bottom_border(cell, size=6)
            for paragraph in cell.paragraphs:
                paragraph.style = "Table Text"
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing_rule = (
                    WD_LINE_SPACING.SINGLE
                )
                for run in paragraph.runs:
                    _set_run_font(
                        run,
                        BODY_CJK_FONT,
                        TABLE_FONT_SIZE_PT,
                        bold=True if row_index == 0 else None,
                    )
        if row_index == 0:
            _repeat_table_header(row)


def _set_table_borders(table) -> None:
    """设置三线表的顶线与底线并移除竖线。"""
    table_properties = table._tbl.tblPr
    existing = table_properties.find(qn("w:tblBorders"))
    if existing is not None:
        table_properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge, value, size in (
        ("top", "single", "12"),
        ("left", "nil", "0"),
        ("bottom", "single", "12"),
        ("right", "nil", "0"),
        ("insideH", "nil", "0"),
        ("insideV", "nil", "0"),
    ):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), value)
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    table_properties.append(borders)


def _set_cell_bottom_border(cell, *, size: int) -> None:
    """设置表头底线。"""
    cell_properties = cell._tc.get_or_add_tcPr()
    borders = cell_properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        cell_properties.append(borders)
    bottom = borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        borders.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "0")
    bottom.set(qn("w:color"), "000000")


def _set_cell_margins(
    cell,
    *,
    top: int,
    start: int,
    bottom: int,
    end: int,
) -> None:
    """设置表格单元格内边距，单位为 twip。"""
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)
    for margin_name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = margins.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _repeat_table_header(row) -> None:
    """让跨页表格重复显示表头。"""
    row_properties = row._tr.get_or_add_trPr()
    header = row_properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        row_properties.append(header)
    header.set(qn("w:val"), "true")


def _add_page_number(paragraph: Paragraph) -> None:
    """在页脚中央插入 PAGE 域。"""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if "PAGE" in paragraph._p.xml:
        return
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    _set_run_font(run, BODY_CJK_FONT, BODY_FONT_SIZE_PT)


def _compact_heading_text(value: str) -> str:
    """移除标题中的空白，兼容“目 录”等写法。"""
    return re.sub(r"\s+", "", value or "")


def _is_level_one_heading(paragraph: Paragraph) -> bool:
    """判断段落是否为 Word 一级标题。"""
    style_name = (paragraph.style.name or "").lower()
    return style_name.startswith("heading 1")


def _insert_toc_field_after(
    document: DocumentObject,
    paragraph: Paragraph,
) -> Paragraph:
    """在目录标题后插入可由 Word 自动更新的三级目录域。"""
    toc_element = OxmlElement("w:p")
    paragraph._p.addnext(toc_element)
    toc_paragraph = Paragraph(toc_element, paragraph._parent)
    toc_paragraph.paragraph_format.first_line_indent = Pt(0)

    run = toc_paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在打开文档时自动更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])
    _set_run_font(run, BODY_CJK_FONT, BODY_FONT_SIZE_PT)

    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    return toc_paragraph


def _prepare_docx_front_matter(document: DocumentObject) -> None:
    """把静态目录替换为 Word 目录域，并确保目录与正文分别起页。"""
    paragraphs = list(document.paragraphs)
    toc_index = next(
        (
            index
            for index, paragraph in enumerate(paragraphs)
            if _compact_heading_text(paragraph.text) == "目录"
        ),
        None,
    )
    if toc_index is None:
        return

    toc_heading = paragraphs[toc_index]
    body_heading_index = next(
        (
            index
            for index in range(toc_index + 1, len(paragraphs))
            if _is_level_one_heading(paragraphs[index])
        ),
        len(paragraphs),
    )
    for paragraph in paragraphs[toc_index + 1 : body_heading_index]:
        element = paragraph._element
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    _insert_toc_field_after(document, toc_heading)
    if body_heading_index < len(paragraphs):
        body_heading = paragraphs[body_heading_index]
        body_heading.paragraph_format.page_break_before = True


def _contains_equation(paragraph: Paragraph) -> bool:
    """判断段落是否为块级 Office Math，排除正文中的行内公式。"""
    if paragraph._p.xpath(".//m:oMathPara"):
        return True
    if not paragraph._p.xpath(".//m:oMath"):
        return False
    outside_math_text = "".join(
        node.text or ""
        for node in paragraph._p.xpath(
            ".//w:t[not(ancestor::m:oMath)]"
        )
    ).strip()
    return not outside_math_text


def _contains_drawing(paragraph: Paragraph) -> bool:
    """判断段落是否包含图片或旧式绘图。"""
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def _is_caption(text: str, style_name: str) -> bool:
    """判断段落是否为图题或表题。"""
    return (
        "caption" in style_name
        or bool(re.match(r"^\s*[图表]\s*\d+", text))
    )


def _looks_like_filename(value: str) -> bool:
    """判断图题是否只是文件名。"""
    candidate = value.strip().strip("`")
    return bool(
        not candidate
        or re.search(
            r"\.(?:png|jpe?g|gif|bmp|webp|svg)(?:[?#].*)?$",
            candidate,
            re.IGNORECASE,
        )
        or "/" in candidate
        or "\\" in candidate
    )


def _looks_like_document_title(value: str) -> bool:
    """判断首个一级标题是否更像论文总标题。"""
    if not value:
        return False
    common_sections = (
        "摘要",
        "目录",
        "问题重述",
        "问题分析",
        "模型假设",
        "符号说明",
        "参考文献",
        "附录",
        "致谢",
    )
    if value.startswith(common_sections):
        return False
    return not bool(
        re.match(
            r"^(?:[一二三四五六七八九十]+、|"
            r"\d+(?:\.\d+)*[.\s、]|第[一二三四五六七八九十\d]+章)",
            value,
        )
    )


def add_page_break(paragraph: Paragraph) -> None:
    """为需要独立成页的前置部分添加分页符。

    Args:
        paragraph: 目标段落。
    """
    paragraph.runs[0].add_break(WD_BREAK.PAGE) if paragraph.runs else paragraph.add_run().add_break(
        WD_BREAK.PAGE
    )
